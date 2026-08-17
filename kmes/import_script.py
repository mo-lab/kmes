import os
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'kmes.settings')
django.setup()

from openpyxl import load_workbook
from core.models import EquipmentTag, Project,Area,System
from django.db import models

def import_excel(file_path, project_id):
	wb = load_workbook(file_path)
	ws = wb.active
	project = Project.objects.get(pk=5)
	area = Area.objects.get(pk=2)
	sys = System.objects.get(pk=3)
	
	for row_idx in range(2, ws.max_row + 1):
		description = ws.cell(row=row_idx, column=4).value
		remarks = ws.cell(row=row_idx, column=5).value
		drawing = ws.cell(row=row_idx, column=7).value
		
		tag = EquipmentTag(
				tag_number=str(remarks) if remarks else '',
				project=project,
				status='ENG',
				area=area,
				system=sys,
				description=description if description else '',
				drawing_num=drawing if drawing else '',
				)
		tag.save()
		print(f'Created: {tag.tag_number}')


if __name__ == '__main__':
	import_excel('p1.xlsx', 1)
