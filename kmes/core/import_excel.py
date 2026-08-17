from openpyxl import load_workbook
from django.core.management.base import BaseCommand
from core.models import EquipmentTag, Project

class Command(BaseCommand):
	help = 'Import equipment tags from Excel file using openpyxl'
	
	def add_arguments(self, parser):
		parser.add_argument('file_path', type=str, help='Path to Excel file')
		parser.add_argument('--project_id', type=int, help='Project ID')
		parser.add_argument('--start_row', type=int, default=2, help='Start row (1-indexed, default: 2)')
	
	def handle(self, *args, **options):
		file_path = options['file_path']
		project_id = options.get('project_id')
		start_row = options.get('start_row', 2)
		
		self.stdout.write(f'Reading Excel file: {file_path}')
		
		try:
			# بارگذاری فایل
			wb = load_workbook(file_path, data_only=True)
			ws = wb.active
			
			# پیدا کردن ستون‌ها بر اساس سطر اول (header)
			headers = {}
			for col_idx, cell in enumerate(ws[1], 1):
				if cell.value:
					header = str(cell.value).lower().strip()
					if 'description' in header or 'توضیحات' in header:
						headers['description'] = col_idx
					elif 'remark' in header or 'نکات' in header:
						headers['remarks'] = col_idx
					elif 'drawing' in header or 'نقشه' in header:
						headers['drawing'] = col_idx
			
			self.stdout.write(f'Found columns: {headers}')
			
			if 'description' not in headers:
				self.stderr.write('Error: Could not find description column')
				return
			
			# دریافت پروژه
			project = None
			if project_id:
				try:
					project = Project.objects.get(pk=project_id)
					self.stdout.write(f'Using project: {project.name}')
				except Project.DoesNotExist:
					self.stderr.write(f'Project with ID {project_id} not found!')
					return
			
			# پردازش داده‌ها
			success_count = 0
			error_count = 0
			
			for row_idx in range(start_row, ws.max_row + 1):
				try:
					# خواندن مقادیر
					description = ws.cell(row=row_idx, column=headers['description']).value or ''
					remarks = ''
					drawing = ''
					
					if 'remarks' in headers:
						remarks = ws.cell(row=row_idx, column=headers['remarks']).value or ''
					
					if 'drawing' in headers:
						drawing = ws.cell(row=row_idx, column=headers['drawing']).value or ''
					
					# ایجاد برچسب
					tag_number = f"TAG-{row_idx - start_row + 1:04d}"
					
					equipment_tag = EquipmentTag(
							tag_number=tag_number,
							description=str(description),
							project=project,
							status='ENG',
							equipment_type='OTHR',
							)
					
					# اضافه کردن notes
					notes_parts = []
					if remarks:
						notes_parts.append(f"Remarks: {remarks}")
					if drawing:
						notes_parts.append(f"Drawing: {drawing}")
					
					if notes_parts:
						equipment_tag.notes = '\n'.join(notes_parts)
					
					equipment_tag.save()
					success_count += 1
				
				except Exception as e:
					error_count += 1
					self.stderr.write(f'Error at row {row_idx}: {e}')
			
			# نمایش نتیجه
			self.stdout.write(f'\n✅ Successfully imported: {success_count}')
			self.stdout.write(f'❌ Errors: {error_count}')
		
		except Exception as e:
			self.stderr.write(f'Error: {e}')
			raise